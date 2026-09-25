"""Render the Swift v3 test report into the supplied RUP Word template.

Keeps the template's own cover, styles, headers/footers, section properties and
table borders; replaces the body with content parsed from the Markdown source.
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from docx.shared import Pt

ROOT = Path(r"C:\Users\ASUS\Desktop\Swif Shazan\Swift")
TEMPLATE = ROOT / "Testing" / "Template for Test plan.docx"
SOURCE = ROOT / "Testing" / "Swift_Master_Test_Plan_and_Report_v3.md"
OUTPUT = ROOT / "Testing" / "Swift_Master_Test_Plan_and_Report_v3.docx"

COVER_REPLACEMENTS = {
    "<Project Name>": "Swift",
    "<Iteration/ Master> Test Plan": "Master Test Plan and Test Report",
    "Version <1.0>": "Version 3.0",
    "<Company Name>": "Swift",
    "2024": "2026",
}


# --------------------------------------------------------------------------- xml helpers

def replace_text(element, replacements: dict[str, str]) -> None:
    for node in element.xpath(".//w:t"):
        value = node.text or ""
        for old, new in replacements.items():
            value = value.replace(old, new)
        node.text = value


def suppress_heading_number(paragraph) -> None:
    """The headings carry their own numbers, so switch off list numbering."""
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        num_id = OxmlElement("w:numId")
        num_pr.append(num_id)
    num_id.set(qn("w:val"), "0")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_update_fields(document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field to build the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


# --------------------------------------------------------------------------- inline runs

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|<https?://[^>]+>)")


def add_runs(paragraph, text: str) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif piece.startswith("<http"):
            paragraph.add_run(piece[1:-1])
        else:
            paragraph.add_run(piece)


# --------------------------------------------------------------------------- markdown parse

def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_blocks(lines: list[str]):
    """Yield (kind, payload) blocks: heading, para, bullet, number, table, code, rule."""
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            body: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                body.append(lines[index].rstrip())
                index += 1
            index += 1
            yield ("code", body)
            continue

        if set(stripped) <= {"-"} and len(stripped) >= 3:
            yield ("rule", None)
            index += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            yield ("heading", (level, stripped.lstrip("# ").strip()))
            index += 1
            continue

        if stripped.startswith("|"):
            table: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = split_row(lines[index])
                if not all(set(cell) <= set("-: ") and cell for cell in cells):
                    table.append(cells)
                index += 1
            yield ("table", table)
            continue

        if stripped.startswith("- "):
            items: list[str] = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(lines[index].strip()[2:])
                index += 1
            yield ("bullet", items)
            continue

        if re.match(r"^\d+\.\s", stripped):
            items = []
            while index < len(lines) and re.match(r"^\d+\.\s", lines[index].strip()):
                items.append(re.sub(r"^\d+\.\s", "", lines[index].strip()))
                index += 1
            yield ("number", items)
            continue

        yield ("para", stripped)
        index += 1


# --------------------------------------------------------------------------- build

def main() -> None:
    document = Document(TEMPLATE)
    body = document.element.body

    template_children = list(body)
    cover = [deepcopy(node) for node in template_children[:6]]
    section_properties = deepcopy(body.sectPr)
    thin_table_pr = deepcopy(document.tables[0]._tbl.tblPr)
    thin_table_grid = document.tables[0]._tbl.tblGrid
    technique_table_pr = deepcopy(document.tables[1]._tbl.tblPr)

    for node in list(body):
        body.remove(node)
    for node in cover:
        replace_text(node, COVER_REPLACEMENTS)
        body.append(node)
    # python-docx measures column widths from the trailing sectPr and inserts
    # new paragraphs and tables before it, so it has to go back now.
    body.append(section_properties)

    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()

    # The cover already carries title and version; drop the Markdown's own cover lines.
    start = next(i for i, line in enumerate(lines) if line.strip() == "### Revision History")
    lines = lines[start:]

    first_table_seen = False
    toc_inserted = False

    for kind, payload in parse_blocks(lines):
        if kind == "rule":
            continue

        if kind == "heading":
            level, text = payload
            if text == "Revision History":
                paragraph = document.add_paragraph(text, style="Title")
                continue
            style = f"Heading {min(level, 4)}"
            paragraph = document.add_paragraph(style=style)
            add_runs(paragraph, text)
            suppress_heading_number(paragraph)
            continue

        if kind == "para":
            paragraph = document.add_paragraph(style="Body Text")
            add_runs(paragraph, payload)
            continue

        if kind == "code":
            for line in payload:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(line)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                paragraph.paragraph_format.space_after = Pt(0)
            continue

        if kind in {"bullet", "number"}:
            for position, item in enumerate(payload, 1):
                prefix = "\u2022  " if kind == "bullet" else f"{position}.  "
                paragraph = document.add_paragraph(style="Body Text")
                paragraph.add_run(prefix)
                add_runs(paragraph, item)
                paragraph.paragraph_format.left_indent = Pt(18)
                paragraph.paragraph_format.space_after = Pt(2)
            continue

        if kind == "table":
            rows = payload
            # A technique table opens with "| | |", which is a spacer, not a header.
            while rows and not any(cell.strip() for cell in rows[0]):
                rows = rows[1:]
            if not rows:
                continue
            width = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=width)
            # Two-column tables are the template's technique tables; the rest are thin grids.
            source_pr = technique_table_pr if width == 2 else thin_table_pr
            table._tbl.replace(table._tbl.tblPr, deepcopy(source_pr))

            header_is_blank = width == 2  # technique tables are label/value, not headed
            for r, row in enumerate(rows):
                for c in range(width):
                    cell = table.cell(r, c)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    text = row[c] if c < len(row) else ""
                    add_runs(paragraph, text)
                    paragraph.paragraph_format.space_after = Pt(2)
                    if r == 0 and not header_is_blank:
                        for run in paragraph.runs:
                            run.bold = True
            if not header_is_blank:
                set_repeat_table_header(table.rows[0])

            if not first_table_seen and width == 4:
                table._tbl.replace(table._tbl.tblGrid, deepcopy(thin_table_grid))
                first_table_seen = True
                # The template puts the contents page straight after revision history.
                if not toc_inserted:
                    document.add_paragraph("Table of Contents", style="Title")
                    add_toc_field(document.add_paragraph())
                    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    toc_inserted = True
            document.add_paragraph()
            continue

    for part in document.part.package.parts:
        name = str(part.partname)
        if "/header" in name or "/footer" in name:
            element = getattr(part, "element", None)
            if element is not None:
                replace_text(element, COVER_REPLACEMENTS)

    document.core_properties.title = "Master Test Plan and Test Report"
    document.core_properties.subject = "Swift"
    document.core_properties.author = "Swift project team"
    document.core_properties.keywords = (
        "testing, test plan, test report, evidence audit, data science, error analysis"
    )
    set_update_fields(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
