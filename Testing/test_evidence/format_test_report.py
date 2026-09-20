"""Apply the supplied RUP Word template structure to the Swift test report."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Testing" / "Template for Test plan.docx"
DEFAULT_OUTPUT = ROOT / "Testing" / "Swift_Master_Test_Plan_and_Report.docx"


def element_text(element) -> str:
    return "".join((node.text or "") for node in element.xpath(".//w:t"))


def replace_text(element, replacements: dict[str, str]) -> None:
    for node in element.xpath(".//w:t"):
        value = node.text or ""
        for old, new in replacements.items():
            value = value.replace(old, new)
        node.text = value


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def suppress_heading_number(paragraph_element) -> None:
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        paragraph_element.insert(0, p_pr)
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        num_id = OxmlElement("w:numId")
        num_pr.append(num_id)
    num_id.set(qn("w:val"), "0")


def update_extended_properties(path: Path) -> None:
    temp = path.with_suffix(".tmp.docx")
    app_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    with ZipFile(path, "r") as source, ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "docProps/app.xml":
                root = etree.fromstring(data)
                company = root.find(f"{{{app_namespace}}}Company")
                if company is None:
                    company = etree.SubElement(root, f"{{{app_namespace}}}Company")
                company.text = "Swift"
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(item, data)
    temp.replace(path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Apply the supplied RUP Word template to a generated report."
    )
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    template = Document(TEMPLATE)
    report = Document(args.source)
    body = report.element.body

    # Remove Pandoc's generic title block. The template-native cover is inserted below.
    for child in list(body)[:4]:
        body.remove(child)

    # Move the unnumbered revision history before the generated TOC, as in the template.
    children = list(body)
    revision_heading = next(
        child for child in children
        if child.tag == qn("w:p") and "Revision History" in element_text(child)
    )
    revision_index = children.index(revision_heading)
    cluster_start = revision_index - 1 if revision_index and children[revision_index - 1].tag == qn("w:bookmarkStart") else revision_index
    cluster_end = revision_index + 1
    while cluster_end + 1 < len(children) and children[cluster_end + 1].tag in {
        qn("w:tbl"), qn("w:bookmarkEnd")
    }:
        cluster_end += 1
    revision_cluster = children[cluster_start:cluster_end + 1]
    for child in revision_cluster:
        body.remove(child)

    # Match the template's unnumbered Title treatment for Revision History.
    p_pr = revision_heading.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        revision_heading.insert(0, p_pr)
    p_style = p_pr.find(qn("w:pStyle"))
    if p_style is None:
        p_style = OxmlElement("w:pStyle")
        p_pr.insert(0, p_style)
    p_style.set(qn("w:val"), "Title")

    for paragraph in body.findall(qn("w:p")):
        if element_text(paragraph).strip() in {"Executive Test Summary", "Release recommendation"}:
            suppress_heading_number(paragraph)

    # Copy the original cover nodes and its section break verbatim.
    template_children = list(template.element.body)
    cover = [deepcopy(node) for node in template_children[:6]]
    cover_break = deepcopy(template_children[7])
    break_style = cover_break.find("./w:pPr/w:pStyle", namespaces=cover_break.nsmap)
    if break_style is not None:
        break_style.set(qn("w:val"), "Normal")

    replacements = {
        "<Project Name>": "Swift",
        "<Iteration/ Master> Test Plan": "Master Test Plan and Test Report",
        "Version <1.0>": "Version 2.0",
        "<Company Name>": "Swift",
        "2024": "2026",
    }
    for node in [*cover, cover_break]:
        replace_text(node, replacements)

    insertion = 0
    for node in [*cover, cover_break, *revision_cluster]:
        body.insert(insertion, node)
        insertion += 1

    # Use the template's normal body section (including its distinct top margin
    # and header/footer references) instead of the vertically centred cover section.
    current_final_section = body.sectPr
    body.replace(current_final_section, deepcopy(template.element.body.sectPr))

    # Apply the actual template table borders/layout, not Pandoc's default table style.
    thin_table_properties = template.tables[0]._tbl.tblPr
    technique_table_properties = template.tables[1]._tbl.tblPr
    for table in report.tables:
        source_properties = technique_table_properties if len(table.columns) == 2 else thin_table_properties
        table._tbl.replace(table._tbl.tblPr, deepcopy(source_properties))
        if table.rows:
            set_repeat_table_header(table.rows[0])

    # The revision-history table uses the template's exact four-column proportions.
    revision_table = next(table for table in report.tables if table.cell(0, 0).text.strip() == "Date")
    revision_table._tbl.replace(revision_table._tbl.tblGrid, deepcopy(template.tables[0]._tbl.tblGrid))

    # Replace cached field values in every copied header/footer part.
    for part in report.part.package.parts:
        if "/header" in str(part.partname) or "/footer" in str(part.partname):
            element = getattr(part, "element", None)
            if element is not None:
                replace_text(element, replacements)

    report.core_properties.title = "Master Test Plan and Test Report"
    report.core_properties.subject = "Swift"
    report.core_properties.author = "Swift project team / Codex test execution"
    report.core_properties.keywords = "testing, test plan, test report, data science, error analysis"
    set_update_fields(report)
    report.save(args.output)
    update_extended_properties(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
